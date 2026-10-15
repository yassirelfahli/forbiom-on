"""Build the Activity 3 ABM simulation report (Word) from ./results, ./figures and ./calibration."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from abm import BASE_PARAMS, DIST_GTA, REGIONS, STRATEGIES, YEARS, load_calibration
from experiments import CRIT_LABEL, DEMAND, SENSITIVITY, WEIGHTS

HERE = Path(__file__).resolve().parent
RES, TAB, FIG = HERE / "results", HERE / "results" / "tables", HERE / "figures"
OUT = HERE / "Activity3_ABM_Simulation_Report.docx"

REG = {"TroutLake": "Trout Lake", "Abitibi": "Abitibi River", "Algoma": "Algoma", "Kenora": "Kenora",
       "Nipissing": "Nipissing", "Algonquin": "Algonquin Park"}
STRATS = ["S0", "S1", "S2", "S3", "S4", "S5", "S6"]
SNAME = {k: v["name"] for k, v in STRATEGIES.items()} | {"S6": "Regionally differentiated portfolio"}
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


# ---------------------------------------------------------------------------------------------
# docx helpers
# ---------------------------------------------------------------------------------------------
class Report:
    def __init__(self) -> None:
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        st = self.doc.styles
        st["Normal"].font.name = "Calibri"
        st["Normal"].font.size = Pt(10.5)
        st["Normal"].paragraph_format.space_after = Pt(6)
        st["Normal"].paragraph_format.line_spacing = 1.12
        for lvl, size in ((1, 15), (2, 12.5), (3, 11)):
            h = st[f"Heading {lvl}"]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.color.rgb = NAVY
        sec = self.doc.sections[0]
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        self._footer(sec)

    def _footer(self, sec) -> None:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("FORBIOM-ABM | Activity 3 ABM Simulation Report | page ")
        r.font.size = Pt(8)
        run = p.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)
        run.font.size = Pt(8)

    def h(self, text: str, level: int = 1) -> None:
        self.doc.add_heading(text, level=level)

    def p(self, text: str = "", bold_lead: str | None = None, italic: bool = False, size: float | None = None,
          align=None) -> None:
        para = self.doc.add_paragraph()
        if bold_lead:
            para.add_run(bold_lead).bold = True
        r = para.add_run(text)
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if align:
            para.alignment = align

    def bullets(self, items: list[str], style: str = "List Bullet") -> None:
        for it in items:
            para = self.doc.add_paragraph(style=style)
            if it.startswith("**") and "**" in it[2:]:
                lead, rest = it[2:].split("**", 1)
                para.add_run(lead).bold = True
                para.add_run(rest)
            else:
                para.add_run(it)

    def figure(self, name: str, caption: str, width_cm: float = 16.0) -> None:
        self.fig_n += 1
        self.doc.add_picture(str(FIG / f"{name}.png"), width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = self.doc.add_paragraph()
        r = c.add_run(f"Figure {self.fig_n}. ")
        r.bold = True
        r.font.size = Pt(9)
        r2 = c.add_run(caption)
        r2.font.size = Pt(9)
        r2.italic = True

    def table(self, df: pd.DataFrame, caption: str, col_widths: list[float] | None = None,
              font: float = 8.5, note: str | None = None) -> None:
        self.tab_n += 1
        c = self.doc.add_paragraph()
        r = c.add_run(f"Table {self.tab_n}. ")
        r.bold = True
        r.font.size = Pt(9)
        r2 = c.add_run(caption)
        r2.font.size = Pt(9)
        t = self.doc.add_table(rows=1, cols=len(df.columns))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, col in enumerate(df.columns):
            cell = t.rows[0].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(col))
            run.bold = True
            run.font.size = Pt(font)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "DCE6F1")
            cell._tc.get_or_add_tcPr().append(shd)
        for _, row in df.iterrows():
            cells = t.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = ""
                run = cells[j].paragraphs[0].add_run(str(v))
                run.font.size = Pt(font)
        if col_widths:
            for row in t.rows:
                for j, w in enumerate(col_widths):
                    row.cells[j].width = Cm(w)
        if note:
            n = self.doc.add_paragraph()
            rn = n.add_run(note)
            rn.font.size = Pt(8)
            rn.italic = True
        else:
            self.doc.add_paragraph()

    def page_break(self) -> None:
        self.doc.add_page_break()


def f0(x): return f"{x:,.0f}"
def f1(x): return f"{x:,.1f}"
def f2(x): return f"{x:,.2f}"
def pct(x): return f"{100 * x:.0f}%"


# ---------------------------------------------------------------------------------------------
# content
# ---------------------------------------------------------------------------------------------
def main() -> None:
    cal = load_calibration()
    prov = pd.read_csv(TAB / "provincial_base.csv", header=[0, 1], index_col=0)
    M = prov.xs("mean", axis=1, level=1)
    Q05 = prov.xs("q05", axis=1, level=1)
    Q95 = prov.xs("q95", axis=1, level=1)
    reg = pd.read_csv(TAB / "regional_base.csv").set_index(["strategy", "region"])
    calib = pd.read_csv(TAB / "calibration_check.csv", index_col=0)
    by_dem = pd.read_csv(TAB / "provincial_by_demand.csv").set_index(["scenario", "strategy"])
    sens = pd.read_csv(TAB / "sensitivity_ontario.csv").set_index(["scenario", "strategy"])
    plants = pd.read_csv(TAB / "plants_base.csv").set_index(["strategy", "tech", "region"])
    variants = pd.read_csv(RES / "regional_variants.csv")
    rob = pd.read_csv(RES / "rank_robustness.csv", index_col=0)
    pscore = pd.read_csv(RES / "provincial_scores_Base.csv", index_col=0)
    s6def = json.loads((RES / "s6_definition.json").read_text())
    jobs_efar = pd.read_csv(HERE / "calibration" / "efar_jobs.csv")
    P = BASE_PARAMS
    runs = pd.read_csv(RES / "runs_summary.csv")
    ont = runs[runs.region == "Ontario"]
    n_runs = ont.query("scenario=='Base' and strategy=='S0'").shape[0]
    n_sims = len(ont) + (pd.read_csv(RES / "sensitivity_summary.csv").region == "Ontario").sum()
    equity = ont.query("scenario=='Base'").groupby("strategy").ind_equity_income_cum.mean()
    hist = pd.read_csv(HERE.parent / "arwoods_agg_by_year.csv", dtype=str)
    hist = hist[hist.FiscalYear.astype(int).between(2010, 2024)]
    hist_util = 1 - hist.Share_gap_over_LTMD.str.replace(",", ".").astype(float)

    act_tot, ltmd_tot = cal.ACT_total.sum(), cal.LTMD_total.sum()
    gap_tot = ltmd_tot - act_tot
    gap_sw = (cal.LTMD_SW - cal.ACT_SW).sum()
    gap_po = (cal.LTMD_PO - cal.ACT_PO).sum()
    gap_hw = (cal.LTMD_HW - cal.ACT_HW).sum()
    s0, s6 = M.loc["S0"], M.loc["S6"]
    add_harv = s6.harvest_cum_Mm3 - s0.harvest_cum_Mm3
    va_gain = s6.va_npv - s0.va_npv
    leverage = va_gain / s6.public_cost
    north_units_cum = sum(P["north_housing_units"].values()) * sum((1 + P["north_housing_growth"]) ** t for t in range(len(YEARS)))
    panel_units_local = s6.panel_local_cum_k * 1000 / P["m3_per_unit"]
    po_late = sum(reg.loc[("S6", r), "harvest_PO_late"] for r in REGIONS)
    hw_late = sum(reg.loc[("S6", r), "harvest_HW_late"] for r in REGIONS)
    max_err = calib.error_pct.abs().max()

    R = Report()
    d = R.doc

    # ---- title page ------------------------------------------------------------------------
    for _ in range(3):
        d.add_paragraph()
    R.p("ACTIVITY 3 DELIVERABLE: ABM SIMULATION REPORT", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = t.add_run("Agent-Based Modelling of Forest Biomass Utilization Strategies across Ontario's Regions, "
                   "with a Focus on Indigenous and Rural Communities")
    tr.bold = True
    tr.font.size = Pt(20)
    tr.font.color.rgb = NAVY
    R.p("Results from FORBIOM-ABM, the agent-based companion to the FORBIOM-ON forest value-chain model",
        italic=True, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    d.add_paragraph()
    R.p("Project: Optimizing Biomass-Based Modular Mass Timber Housing for Affordable Construction in Ontario "
        "Using Agent-Based Modeling", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    R.p("Ontario Transfer Payment Agreement, Ministry of Natural Resources (Schedule \"C\", Activity 3)",
        size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(3):
        d.add_paragraph()
    R.p("Forest Economy and Decision Support Systems (FEDSS) Lab\nBioeconomy and Sustainable Materials Institute (BSMI)\n"
        "Institute of Forestry and Conservation, University of Toronto", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    R.p("Rasoul Yousefpour, Ph.D. (Principal Investigator)\nYasser Elfahli, M.A.", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    R.p("October 2026", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(4):
        d.add_paragraph()
    R.p("Acknowledgement and disclaimer. This project is supported by the Government of Ontario through the Ministry of "
        "Natural Resources. The views expressed in this report are the views of the authors and do not necessarily "
        "reflect those of the Province.", italic=True, size=9)
    R.page_break()

    # ---- executive summary -----------------------------------------------------------------
    R.h("Executive summary")
    R.p(f"The six Forest Management Units (FMUs) represented in FORBIOM-ON (Trout Lake, Abitibi River, Algoma, Kenora, "
        f"Nipissing and Algonquin Park) harvested on average {act_tot / 1e6:.2f} million m3 per year in 2019-2024, against a "
        f"long-term managed directional volume (LTMD) of {ltmd_tot / 1e6:.2f} million m3 per year (ARWoodS 2025). About "
        f"{gap_tot / 1e6:.2f} million m3 per year ({pct(gap_tot / ltmd_tot)}) of sustainably available fibre is therefore unused, "
        f"of which {gap_sw / 1e6:.2f} million m3 is softwood, {gap_po / 1e6:.2f} million m3 poplar and {gap_hw / 1e6:.2f} million m3 "
        "hardwood. Utilization ranges from 15% in Kenora and 20% in Algoma to 69% in Trout Lake.")
    R.p("Activity 3 developed and applied FORBIOM-ABM, an agent-based model in which FMU managers, harvest contractors "
        "(Indigenous-owned and other), incumbent mills, prospective investors in new processing capacity, product and heat "
        "markets, and a policy agent interact every year from 2026 to 2045. Unlike the optimization in FORBIOM-ON (Activity 2), "
        "the ABM does not assume a central planner: investment, harvest and exit decisions emerge from heterogeneous agents "
        "with bounded rationality, adaptive price expectations and imperfect foresight. The model is calibrated to observed "
        f"FMU harvest by species group (start-year error below {max_err:.1f}% in every FMU) and to eFAR employment intensities. "
        f"Seven biomass utilization strategies were simulated with {n_runs} Monte Carlo replications each, under three "
        "mass-timber demand trajectories and six sensitivity cases, and compared with a seven-criterion multi-criteria "
        "decision analysis (MCDA).")
    R.p("Key findings", bold_lead="")
    R.bullets([
        f"**The optimal strategy is a regionally differentiated hub-and-spoke portfolio (S6) backed by an enabling policy "
        f"package. **It ranks first under base, low and high demand, under all six sensitivity cases, and under all four "
        f"MCDA weightings (economic, community, climate and equal priority). It is practically tied with the province-wide "
        f"version of the same policy package (S5) but reaches the same outcomes with fewer plants, lower public cost and lower "
        f"investment risk. Relative to business as usual it raises fibre "
        f"utilization in 2041-45 from {pct(s0.util_late)} to {pct(s6.util_late)} (90% interval {pct(Q05.loc['S6', 'util_late'])}-"
        f"{pct(Q95.loc['S6', 'util_late'])}), mobilizing an additional {add_harv:.1f} million m3 over 2026-2045.",
        f"**The portfolio has four components: **(i) one or two integrated log-to-mass-timber hubs per FMU, which add the "
        "sawmilling capacity that unlocks softwood; (ii) community-scale log-to-panel plants (\"spokes\") with majority "
        "Indigenous ownership serving northern housing and southern prefab markets; (iii) community bioheat plants using mill and "
        "forest residues; and (iv) pellet plants in Abitibi River and Trout Lake only.",
        f"**Indigenous and rural outcomes improve most under this portfolio. **Direct rural and northern jobs rise from "
        f"{f0(s0.rural_jobs_late)} to {f0(s6.rural_jobs_late)} FTE, Indigenous direct jobs from {f0(s0.ind_jobs_late)} to "
        f"{f0(s6.ind_jobs_late)} FTE, and Indigenous community income (wages, contractor margins and equity returns, NPV at 4%) "
        f"from ${f0(s0.ind_income_npv)} M to ${f0(s6.ind_income_npv)} M.",
        f"**The public cost has high leverage. **The portfolio costs about ${f0(s6.public_cost)} M in capital grants and "
        f"capacity-building programs over 20 years and adds ${f0(va_gain)} M in value added (NPV), about ${leverage:.1f} per "
        f"public dollar. The net carbon benefit rises from {f1(s0.carbon_cum)} to {f1(s6.carbon_cum)} MtCO2e.",
        f"**Enabling policy matters as much as technology. **Adding Indigenous fibre-allocation priority, contractor "
        f"capacity-building, public procurement of mass timber and panels, and a clean-heat incentive (S4 to S5) raises "
        f"Indigenous community income by ${f0(M.loc['S5', 'ind_income_npv'] - M.loc['S4', 'ind_income_npv'])} M. It also cuts "
        f"the share of new plants that lose money from {pct(M.loc['S4', 'risk_unprofitable'])} to "
        f"{pct(M.loc['S5', 'risk_unprofitable'])}, for less than $5 M of extra public cost.",
        f"**Mass-timber plants near Toronto (S1) add value added (+${f0(M.loc['S1', 'va_npv'] - s0.va_npv)} M) but do nothing for "
        "biomass utilization, rural jobs or Indigenous communities, **because they process lumber the existing mills already "
        "produce. Northern hubs alone (S2) mobilize fibre but deliver half the Indigenous benefit of the hybrid. The community "
        "network alone (S3) delivers substantial Indigenous benefit at low public cost but leaves most softwood unused.",
        f"**Poplar and hardwood remain underused under every strategy. **In 2041-45 the portfolio harvests about "
        f"{po_late / 1e3:.0f} thousand m3/yr of poplar and {hw_late / 1e3:.0f} thousand m3/yr of hardwood, against LTMDs of "
        f"{cal.LTMD_PO.sum() / 1e3:.0f} and {cal.LTMD_HW.sum() / 1e3:.0f} thousand m3/yr. Closing this gap requires an "
        "additional product pathway, such as poplar OSB/LVL or hardwood glulam/CLT, which Activity 6 should assess.",
    ])
    R.p("Region-specific optimal strategies (Section 6) differ in the scale and mix of technologies the agents adopt:")
    R.bullets([
        "**Abitibi River: **anchor hub region (two hubs, three spokes, bioheat, one pellet plant). It has the largest softwood "
        "surplus and the largest absolute gains.",
        "**Nipissing and Algoma: **early hub plus a full spoke network. These regions are close to southern markets, and "
        "utilization roughly doubles from a low base.",
        "**Trout Lake: **hub plus residue-based bioheat and pellets. Remote communities give the highest bioheat value; spokes "
        "are sized to northern housing demand.",
        "**Kenora: **the largest relative gain (15% to about 50% utilization) and the highest Indigenous share of benefits, but also "
        "the highest investment risk because of the distance to southern markets. Phase in Indigenous-led spokes tied to "
        "regional and Prairie housing demand, and make the hub conditional on secured offtake.",
        "**Algonquin Park: **one hub and a smaller spoke network. Its hardwood-dominated supply needs a hardwood product "
        "pathway to go further.",
    ])
    R.page_break()

    # ---- 1 introduction ------------------------------------------------------------------
    R.h("1. Introduction and objectives")
    R.p("Ontario's Crown forests can sustainably supply far more fibre than is currently harvested. In the six FMUs used as "
        "the study regions of this project, about half of the long-term managed volume is unused, and the shortfall is "
        "greatest in Kenora, Algoma and Nipissing, where processing capacity has contracted since the mid-2000s. At the same time, Ontario's housing targets and growing interest in prefabricated and "
        "panelized construction create new demand for mass timber (CLT, glulam) and wood-based building components.")
    R.p("Schedule \"C\" of the Agreement describes Activity 3 as: \"Undertake ABM simulation identifying optimal biomass "
        "utilization strategies across Ontario's regions, with a focus on Indigenous and rural communities\", with the "
        "milestone \"ABM simulation reports identifying resulting optimal biomass utilization strategies\". This report is "
        "that deliverable. It builds directly on:")
    R.bullets([
        "**Activity 1 (supply chain feasibility): **sourcing locations, logistics and material flows. The ABM uses the same "
        "FMU regions and species groups as the supply chain analysis.",
        "**Activity 2 (processing optimization): **FORBIOM-ON, a spatial multi-commodity NLP model in GAMS. The ABM uses the "
        "same six regions, the CLT/glulam lumber-input ratios (1.20 and 1.10 m3 lumber per m3), the harvest-cost basis "
        "($38/m3), the carbon storage and process-emission factors, and the baseline mass-timber demand caps (about 165,000 "
        "m3/yr across the six regions).",
    ])
    R.p("FORBIOM-ON identifies the system-optimal allocation of fibre, processing and trade for a single decision-maker. "
        "Real outcomes depend on many independent actors (licensees, contractors, First Nations development corporations, "
        "municipalities, investors) who decide under uncertainty and with different objectives and capacities. "
        "Agent-based modelling (ABM) is the standard tool for representing such decentralized, adaptive behaviour and for "
        "testing whether strategies that look optimal on paper are reachable, robust and fair in practice. The objectives "
        "of Activity 3 were to:")
    R.bullets([
        "represent the main actors of Ontario's forest biomass value chain as heterogeneous, interacting agents;",
        "calibrate the model to observed harvest, fibre supply and employment in the six FMU regions;",
        "simulate alternative biomass utilization strategies, combining technology options and policy instruments, under "
        "market uncertainty;",
        "evaluate the strategies against economic, social (rural and Indigenous), environmental and fiscal criteria; and",
        "identify the optimal strategy for each region, with explicit attention to Indigenous and rural communities.",
    ], style="List Number")

    # ---- 2 model description ---------------------------------------------------------------
    R.h("2. Model description (ODD protocol)")
    R.p("The model is described following the Overview, Design concepts and Details (ODD) protocol for agent-based "
        "models (Grimm et al., 2020). The full source code is provided with this report (Appendix B).")
    R.h("2.1 Purpose and patterns", 2)
    R.p("FORBIOM-ABM evaluates how alternative combinations of processing technologies and policies change biomass "
        "utilization, regional economic activity, Indigenous and rural benefits, carbon outcomes and investment risk across "
        "six Ontario FMU regions. The model is designed to reproduce three observed patterns: (i) current FMU-level "
        "harvest well below LTMD, with large differences between regions; (ii) persistent under-utilization of poplar and "
        "hardwood relative to softwood; and (iii) the cyclical sensitivity of harvest to lumber prices, including mill "
        "curtailment and closure after sustained losses.")
    R.figure("fig01_architecture", "Agent types and interactions in FORBIOM-ABM. Each annual time step runs from "
             "environment and policy signals through fibre and product markets to investment, exit and outcome accounting.")
    R.h("2.2 Entities, state variables and scales", 2)
    ent = pd.DataFrame([
        ["FMU manager (6)", "LTMD by species group (SW, PO, HW); available volume after disturbance; allocation rules",
         "Releases sustainable volume; allocates fibre to licence holders, then entrants (Indigenous-majority entrants first under the priority policy)"],
        ["Harvest contractor (10 per FMU initially)", "Capacity (m3/yr), cost multiplier, Indigenous ownership, sales, margin",
         "Offers volume via a logistic response to log price relative to own cost; expands capacity when fully used and profitable"],
        ["Incumbent mill (3 per FMU x species group)", "Capacity calibrated to observed harvest, cost multiplier, consecutive loss years",
         "Sets utilization from expected margin; exits after 3 loss years (probability 0.4)"],
        ["Mass-timber hub (entrant)", "150,000 m3/yr log intake; 60,000 m3/yr CLT/glulam; capex $150 M; Indigenous equity",
         "Integrated sawmill + CLT/glulam line; sells mass timber to southern Ontario, surplus lumber to commodity markets"],
        ["Urban CLT plant (entrant, GTA)", "60,000 m3/yr; capex $100 M; buys lumber from the north",
         "Converts purchased lumber; does not add harvest"],
        ["Community log-to-panel spoke (entrant)", "25,000 m3/yr logs; 10,000 m3/yr panels; capex $25 M; Indigenous equity",
         "Small sawmill + prefab panel line; serves northern community housing first, then southern prefab demand"],
        ["Community bioheat (entrant)", "5,000 odt/yr; capex $12 M", "Displaces heating oil/propane; residues first, then low-grade roundwood"],
        ["Pellet plant (entrant)", "80,000 odt/yr; capex $50 M", "Export pellets via Thunder Bay or Montreal"],
        ["Markets", "Price paths with AR(1) shocks; demand caps", "Southern mass timber and panels (price-clearing), northern housing, heat demand, commodities"],
        ["Policy agent", "Grants, equity rules, allocation priority, capacity building, procurement, clean-heat incentive", "Defined by the strategy (Section 4)"],
    ], columns=["Agent / entity", "Key state variables", "Behaviour"])
    R.table(ent, "Agents and entities in FORBIOM-ABM.", [3.6, 5.6, 7.4], font=8)
    R.p("Spatial scale: six FMU regions plus a Greater Toronto Area (GTA) market/processing node. Road distances are "
        "approximated from FMU centroids (great-circle distance x 1.3 circuity), e.g. "
        + ", ".join(f"{REG[r]} {DIST_GTA[r]:,.0f} km" for r in REGIONS) + " to the GTA. Temporal scale: annual steps, "
        "2026-2045. Monetary values are 2026 Canadian dollars (real).")
    R.h("2.3 Process overview and scheduling", 2)
    R.bullets([
        "Environment: draw correlated price shocks (lumber, OSB, mass timber, panels, pellets, heating fuel) and FMU disturbance "
        "(fire/blowdown probability 6%/yr, 10-30% volume loss).",
        "FMU managers release available volume by species group.",
        "Processors form desired roundwood intake from expected margins (logistic utilization rule).",
        "Contractors offer volume (logistic supply rule); regional fibre markets allocate volume by priority and update "
        "log prices in proportion to excess demand.",
        "Production, residues and product sales: southern mass-timber and panel prices clear against demand (elasticity 1.5); "
        "spokes serve local housing first; bioheat and pellet plants use mill surplus residues, then forest residues, then "
        "low-grade roundwood.",
        "Accounting: profits after capital charges, jobs, Indigenous jobs and income, carbon, public cost.",
        "Investment: potential entrants evaluate NPV with adaptive expectations, fibre and market headroom, learning-adjusted "
        "capex and grants; build with a logistic probability; plants come online after two years with a 50%/80% ramp-up.",
        "Exit and contractor dynamics; agents update price expectations.",
    ], style="List Number")
    R.h("2.4 Design concepts", 2)
    R.bullets([
        "**Emergence: **regional technology portfolios, log prices, fibre utilization, plant failures and the distribution of "
        "benefits emerge from agent interactions; none are imposed.",
        "**Adaptation and objectives: **firms maximize expected profit under bounded rationality, using logistic (sigmoid) "
        "decision rules that translate margins into utilization, supply or investment probabilities. This is the behavioural "
        "harvest rule proposed in the March 2026 preliminary report, sigma(z) = 1/(1+exp(-k(z-z0))).",
        "**Prediction and learning: **adaptive expectations (smoothing weight 0.3) for prices; learning-by-doing reduces capex by "
        "10% per doubling of plants built.",
        "**Sensing: **agents observe last-period regional log prices, product prices, demand headroom and uncommitted fibre in "
        "their FMU.",
        "**Interaction: **indirect, through regional fibre markets (shared contractors, shared LTMD) and product markets (shared "
        "southern demand); direct, through FMU allocation priority.",
        "**Stochasticity: **price shocks, disturbance, agent heterogeneity and discrete investment/exit decisions. Strategies are "
        "compared with common random numbers (identical environment draws and decision-keyed random numbers).",
        "**Observation: **annual regional indicators and plant-level outcomes (Section 4.3).",
    ])
    R.h("2.5 Key submodels", 2)
    R.p("Contractor supply. Contractor i offers q_i = K_i x e_f x sigma(12 x (w_rf - c_irf)/c_irf + 1), where K_i is capacity, "
        "e_f the share of effort to species group f (proportional to buyer demand), w_rf the regional log price and c_irf the "
        "delivered cost (harvest $38/m3 x regional access factor x agent multiplier, plus in-region haul and stumpage).")
    R.p("Mill utilization. Plant j demands D_jf = Cap_j x sigma(4 x (v_jf - E[w_rf])/E[w_rf] + 1), where v_jf is the plant's "
        "willingness to pay per m3 (net product value after conversion cost and product freight, including residue value).")
    R.p("Fibre market. Allocated volume is min(D, S, available LTMD). Licence holders are served before entrants; under the "
        "Indigenous-priority policy, Indigenous-majority entrants and Indigenous contractors are served first. The log price "
        "adjusts by 30% of the relative excess demand each year, bounded by contractor cost and the highest willingness to pay.")
    R.p("Investment. Expected annual margin (at 90% of capacity, with market-share and fibre-headroom checks) gives "
        "NPV = margin x annuity(8%, 20 yr) - capex x (1 - grant) x learning factor. The build probability is "
        "0.6 x sigma(5 x NPV/net capex - 1), with at most two hubs, three spokes, four bioheat plants and one pellet plant per region.")
    R.p("Indigenous accounting. Indigenous income is the sum of (i) wages of Indigenous workers (regional Indigenous labour "
        "share; 60% local hiring at Indigenous-majority plants; +15 percentage points at partnership plants; 80% at "
        "Indigenous-owned contractors), (ii) margins of Indigenous-owned contractors and (iii) the equity share of plant profits "
        "after capital charges. Equity losses are counted, so Indigenous investors carry downside risk in the model.")
    R.p("Carbon accounting. Net benefit = carbon stored in products (FORBIOM-ON factors) + fossil emissions avoided by "
        "bioheat (1.23 tCO2e/odt replacing heating oil) - process, freight and harvest emissions. Biogenic emissions from "
        "bioheat are treated as carbon-neutral at the point of combustion, following energy-sector reporting conventions. "
        "Full life-cycle and substitution effects are part of Activity 4.")

    # ---- 3 data and calibration -------------------------------------------------------------
    R.h("3. Data and calibration")
    R.p("FMU-level supply and harvest were taken from the ARWoodS 2025 extract (merchantable volume, fiscal years "
        "2019-2024). Species groups were aggregated into softwood (SPF, white/red pine, other conifer, cedar), poplar "
        "(PO) and hardwood (white birch, tolerant hardwoods). Employment intensities come from the eFAR summary (2020-2022): "
        + "; ".join(f"{r.FacilityTypeName.lower()} {r.facility_jobs_per_1000m3:.2f}" for r in jobs_efar.itertuples()
                    if r.FacilityTypeName in ("Sawmill", "Composite", "Pulp", "Bioproduct"))
        + " direct facility jobs per 1,000 m3 of roundwood consumed. Prices, costs and technology parameters are listed with "
        "their basis in Appendix A.")
    cal_tab = pd.DataFrame({
        "FMU": [REG[r] for r in REGIONS],
        "LTMD (1,000 m3/yr)": [f0(cal.loc[r, "LTMD_total"] / 1e3) for r in REGIONS],
        "Harvest 2019-24 (1,000 m3/yr)": [f0(cal.loc[r, "ACT_total"] / 1e3) for r in REGIONS],
        "Utilization": [pct(cal.loc[r, "utilization"]) for r in REGIONS],
        "Unused SW": [f0((cal.loc[r, "LTMD_SW"] - cal.loc[r, "ACT_SW"]) / 1e3) for r in REGIONS],
        "Unused PO": [f0((cal.loc[r, "LTMD_PO"] - cal.loc[r, "ACT_PO"]) / 1e3) for r in REGIONS],
        "Unused HW": [f0((cal.loc[r, "LTMD_HW"] - cal.loc[r, "ACT_HW"]) / 1e3) for r in REGIONS],
        "Simulated 2026 (1,000 m3)": [f0(calib.loc[r, "simulated_2026"] / 1e3) for r in REGIONS],
        "Error": [f"{calib.loc[r, 'error_pct']:+.1f}%" for r in REGIONS],
    })
    R.table(cal_tab, "Fibre supply, observed harvest and calibration check by FMU (ARWoodS 2025; simulated values are "
            f"the mean of {n_runs} business-as-usual runs).", font=8)
    R.figure("fig02_calibration", "(a) Observed vs simulated start-year harvest by FMU. (b) LTMD vs harvest by species "
             "group, 2019-2024: unused volume is concentrated in softwood in Abitibi River, Trout Lake, Kenora and Nipissing, "
             "and in poplar and hardwood in Abitibi River, Algoma, Algonquin Park and Kenora.")
    R.p("Validation. Beyond matching start-year harvest, business-as-usual runs reproduce the observed long-run pattern of "
        f"persistent under-utilization: province-wide harvest has been {pct(hist_util.min())}-{pct(hist_util.max())} of LTMD in "
        f"every year since 2010 (ARWoodS). The simulated 2041-45 business-as-usual mean is {pct(s0.util_late)} with a 90% range "
        f"of {pct(Q05.loc['S0', 'util_late'])}-{pct(Q95.loc['S0', 'util_late'])}, reflecting lumber-price cycles and occasional "
        "mill exits. Behavioural parameters (logistic slopes, adjustment speeds) were set to give plausible responses. For "
        f"example, a sustained 20% lumber price drop lowers business-as-usual utilization from {pct(s0.util_late)} to "
        f"{pct(sens.loc[('Lumber price -20%', 'S0'), 'util_late'])}, consistent with the 2006-2009 contraction. Their influence "
        "was checked through the sensitivity analysis.")

    # ---- 4 strategies -----------------------------------------------------------------------
    R.h("4. Strategies and experimental design")
    R.h("4.1 Biomass utilization strategies", 2)
    strat_tab = pd.DataFrame([
        ["S0", SNAME["S0"], "Incumbent mills only; no new technologies or policies"],
        ["S1", SNAME["S1"], "Up to three 60,000 m3 CLT/glulam plants near the GTA, buying northern lumber; no grants"],
        ["S2", SNAME["S2"], "Integrated log-to-CLT/glulam hubs in FMU regions; 15% capital grant; 20% Indigenous equity partnership"],
        ["S3", SNAME["S3"], "Community log-to-panel spokes and bioheat in all FMUs; 51% Indigenous ownership; 40-50% grants; contractor capacity building"],
        ["S4", SNAME["S4"], "Hubs + spokes + bioheat + pellets (S2 and S3 technologies); grants and equity as above"],
        ["S5", SNAME["S5"], "S4 + Indigenous fibre-allocation priority + contractor capacity building + 25% public procurement of mass timber/panels + $50/tCO2 clean-heat incentive"],
        ["S6", SNAME["S6"], "S5 tuned region by region: technologies dropped where they add cost or risk without benefit (result: pellets only in "
         + " and ".join(REG[r] for r in s6def["strategy"]["techs"]["pellet"]) + ")"],
    ], columns=["Code", "Strategy", "Description"])
    R.table(strat_tab, "Strategies simulated.", [1.2, 4.6, 10.8], font=8)
    R.h("4.2 Scenarios, replication and sensitivity", 2)
    R.p(f"Each strategy was run {n_runs} times (Monte Carlo) under three mass-timber/prefab demand growth trajectories: "
        + ", ".join(f"{k} {v:.0%}/yr" for k, v in DEMAND.items()) + ". Six one-at-a-time sensitivity cases ("
        + "; ".join(SENSITIVITY) + f") were run with 30 replications each. In total {n_sims:,} simulations of 20 years were "
        "performed, plus the regional tuning runs described in Section 4.3 (runtime a few minutes on a laptop).")
    R.h("4.3 Evaluation criteria and MCDA", 2)
    R.p("Seven criteria were evaluated at provincial and FMU level: " + "; ".join(CRIT_LABEL.values()) + ". Value added "
        "and Indigenous income are discounted at a 4% social rate to 2026. Criteria were min-max normalized across "
        "strategies and aggregated with four weighting schemes (" + ", ".join(WEIGHTS) + ") to test the robustness of "
        "rankings to stakeholder priorities. S6 was derived in a second stage. Starting from S5, each technology was removed "
        "(or spokes capped at one) in each FMU in turn, and a change was kept only if it was no worse regionally and "
        "provincially and it reduced public cost (Appendix C).")

    # ---- 5 results --------------------------------------------------------------------------
    R.h("5. Results")
    R.h("5.1 Provincial outcomes", 2)
    rows = []
    for s in STRATS:
        rows.append([s, f"{pct(M.loc[s, 'util_late'])} ({pct(Q05.loc[s, 'util_late'])}-{pct(Q95.loc[s, 'util_late'])})",
                     f1(M.loc[s, "harvest_cum_Mm3"]), f0(M.loc[s, "mt_cum_k"]), f0(M.loc[s, "panel_cum_k"]),
                     f0(M.loc[s, "rural_jobs_late"]), f0(M.loc[s, "ind_jobs_late"]), f0(M.loc[s, "ind_income_npv"]),
                     f0(M.loc[s, "va_npv"]), f1(M.loc[s, "carbon_cum"]), f0(M.loc[s, "public_cost"]),
                     pct(M.loc[s, "risk_unprofitable"])])
    R.table(pd.DataFrame(rows, columns=["Strategy", "Utilization 2041-45 (90% range)", "Harvest 2026-45 (Mm3)",
                                        "Mass timber (1,000 m3)", "Panels (1,000 m3)", "Rural jobs (FTE)",
                                        "Indigenous jobs (FTE)", "Indigenous income NPV ($M)", "Value added NPV ($M)",
                                        "Net carbon (MtCO2e)", "Public cost ($M)", "New plants unprofitable"]),
            f"Provincial outcomes (six FMUs + GTA), base demand, mean of {n_runs} runs. Jobs are 2036-45 averages; cumulative "
            "values cover 2026-2045.", font=7.5,
            note="S0 and S1 have identical rural and Indigenous outcomes because urban plants add no FMU-region activity.")
    R.figure("fig03_utilization_paths", "Provincial fibre utilization (harvest/LTMD) over time. Lines show medians; shaded "
             "bands show the 5-95% Monte Carlo range. S1 coincides with S0 because urban plants do not add harvest.")
    R.figure("fig06_indicators", "Provincial outcomes by strategy (mean and 5-95% interval).")
    R.p(f"Three results stand out. First, adding processing capacity where the fibre is (S2, S4-S6) is what raises "
        f"utilization; the softwood gap is a capacity gap, not a supply gap. Second, community-scale and Indigenous-owned "
        f"assets (S3-S6) are what move Indigenous outcomes: S3 delivers ${f0(M.loc['S3', 'ind_income_npv'])} M of "
        f"Indigenous income with ${f0(M.loc['S3', 'public_cost'])} M of public cost, compared with "
        f"${f0(M.loc['S2', 'ind_income_npv'])} M for S2. Third, the hybrid strategies combine both effects, and the enabling "
        f"policy package (S5, S6) improves every outcome relative to S4, for a marginal increase in public cost, and sharply "
        f"lowers investment risk. With "
        "Indigenous fibre-allocation priority and procurement-backed demand, community plants obtain fibre and offtake more "
        "reliably, so fewer fail.")

    R.h("5.2 Regional utilization and emergent technology portfolios", 2)
    R.figure("fig04_regional_utilization", "Fibre utilization by FMU and strategy (2041-45 mean); the top row shows the "
             "calibrated 2026 value.")
    R.figure("fig05_plants_by_region", "Plants operating in 2045 by FMU under S6 (mean per run).")
    hub = plants.loc["S6"].xs("mt_hub", level="tech") if "mt_hub" in plants.loc["S6"].index.get_level_values(0) else None
    R.p("Under S6, hubs are built in every FMU, but their economics differ sharply. Mean cumulative hub profit over "
        "2026-2045 is "
        + ", ".join(f"{REG[r]} ${f0(hub.loc[r, 'mean_cum_profit_M'])} M" for r in hub.sort_values('mean_cum_profit_M', ascending=False).index)
        + ". The difference is driven by freight to southern markets. In Kenora, "
        f"{pct(hub.loc['Kenora', 'share_unprofitable'])} of hubs lose money over the horizon, and "
        f"{pct(plants.loc[('S6', 'spoke', 'Kenora'), 'share_unprofitable'])} of Kenora spokes do. "
        "Bioheat is profitable everywhere it is built, and most profitable where it displaces expensive heating fuel in remote "
        "communities (Trout Lake, Abitibi River, Kenora). Pellet plants are viable only where residue pools are large "
        "(Abitibi River, Trout Lake).")
    R.figure("fig08_fibre_groups", "Harvest by species group in 2041-45, S0 vs S6, against LTMD.")
    R.p(f"The softwood gap largely closes under S6, but poplar and hardwood mobilization barely changes "
        f"(poplar {po_late / 1e3:.0f} vs LTMD {cal.LTMD_PO.sum() / 1e3:.0f} thousand m3/yr; hardwood {hw_late / 1e3:.0f} vs "
        f"{cal.LTMD_HW.sum() / 1e3:.0f}). Bioheat and pellet plants mostly use residues, which is economically efficient but "
        "leaves low-grade roundwood standing. A dedicated poplar/hardwood value-added pathway is the main missing piece "
        "for Algoma, Algonquin Park, Kenora (poplar) and Abitibi River (poplar).")

    R.h("5.3 Indigenous and rural community outcomes", 2)
    R.figure("fig10_indigenous_regional", "Indigenous direct jobs and Indigenous community income by FMU for four strategies.")
    ind_rows = []
    for r in REGIONS:
        a, b = reg.loc[("S0", r)], reg.loc[("S6", r)]
        ind_rows.append([REG[r], f0(a.rural_jobs_late), f0(b.rural_jobs_late), f0(a.ind_jobs_late), f0(b.ind_jobs_late),
                         f0(a.ind_income_npv), f0(b.ind_income_npv), f0(a.ind_contractor_vol_late / 1e3), f0(b.ind_contractor_vol_late / 1e3),
                         f0(b.panel_local_cum_k * 1000 / P["m3_per_unit"])])
    R.table(pd.DataFrame(ind_rows, columns=["FMU", "Rural jobs S0", "Rural jobs S6", "Indig. jobs S0", "Indig. jobs S6",
                                            "Indig. income S0 ($M NPV)", "Indig. income S6 ($M NPV)",
                                            "Indig. contractor vol. S0 (1,000 m3/yr)", "Indig. contractor vol. S6",
                                            "Northern housing units supplied (S6, 2026-45)"]),
            "Rural and Indigenous outcomes by FMU, business as usual (S0) vs regional portfolio (S6), base demand.", font=7.5)
    R.p(f"Under S6, Indigenous direct employment rises about {s6.ind_jobs_late / s0.ind_jobs_late:.1f}-fold and Indigenous "
        f"community income {s6.ind_income_npv / s0.ind_income_npv:.1f}-fold. The largest absolute gains are in Abitibi River "
        "and Trout Lake; the largest relative gains are in Kenora, Algoma and Nipissing, where Indigenous participation starts "
        f"from a very low base. Equity returns on Indigenous-majority spokes and bioheat plants and on partnership stakes in "
        f"hubs total about ${f0(equity['S6'])} M (undiscounted, 2026-2045). The rest comes from wages and from "
        f"Indigenous-owned harvesting contractors, whose volume rises roughly {sum(reg.loc[('S6', r), 'ind_contractor_vol_late'] for r in REGIONS) / sum(reg.loc[('S0', r), 'ind_contractor_vol_late'] for r in REGIONS):.1f}-fold. "
        f"Spokes supply about {f0(panel_units_local)} housing-unit equivalents of panels to northern communities over "
        f"2026-2045, roughly {pct(panel_units_local / north_units_cum)} of the assumed northern housing need, linking biomass "
        "utilization directly to the project's affordable-housing objective.")

    R.h("5.4 Strategy ranking and robustness", 2)
    R.figure("fig07_mcda", "(a) MCDA scores under four weighting schemes (base demand). (b) Rank under equal weights across "
             "demand scenarios and sensitivity cases (1 = best).")
    R.figure("fig09_demand_scenarios", "Fibre utilization and Indigenous community income under low, base and high "
             "mass-timber demand growth.")
    sens_rows = []
    for sc in SENSITIVITY:
        a, b = sens.loc[(sc, "S0")], sens.loc[(sc, "S6")]
        sens_rows.append([sc, pct(a.util_late), pct(b.util_late), f0(b.va_npv - a.va_npv), f0(b.ind_income_npv - a.ind_income_npv),
                          pct(b.risk_unprofitable), int(rob.loc["S6", sc])])
    R.table(pd.DataFrame(sens_rows, columns=["Sensitivity case", "Utilization S0", "Utilization S6", "Value added gain S6-S0 ($M)",
                                             "Indigenous income gain S6-S0 ($M)", "S6 plants unprofitable", "S6 rank"]),
            "Sensitivity of S6 outcomes (30 runs per case, base demand).", font=8)
    R.p(f"S6 ranks first in every demand scenario and sensitivity case, and S5 second. The two are close: S6 equal-weight "
        f"score {pscore.loc['S6', 'score_Equal']:.3f} vs {pscore.loc['S5', 'score_Equal']:.3f}. S6's advantage is that it "
        "removes pellet plants that add little in four FMUs, which lowers public cost and the share of failing plants without "
        "measurable loss of benefits. Weaker mass-timber demand (Low) "
        f"reduces S6 utilization to {pct(by_dem.loc[('Low', 'S6'), 'util_late'])}, and the lumber-price downturn case "
        f"lowers utilization in all strategies (S0 falls to {pct(sens.loc[('Lumber price -20%', 'S0'), 'util_late'])}). "
        f"In that case S6 value added is still {sens.loc[('Lumber price -20%', 'S6'), 'va_npv'] / sens.loc[('Lumber price -20%', 'S0'), 'va_npv']:.1f} "
        "times business as usual, because mass timber and panels "
        "buffer lumber-price risk. Higher capex (+30%) and a lower mass-timber price (-20%) raise the share of failing "
        "plants to about 6-8%, which is the main downside risk to manage.")

    # ---- 6 optimal strategies ----------------------------------------------------------------
    R.h("6. Optimal biomass utilization strategies by region")
    R.p("Combining the provincial ranking, the regional outcomes and the region-by-region tuning test (Appendix C), the "
        "simulations identify the following optimal strategies. Plant counts are model means for 2045 under base demand and "
        "indicate scale rather than exact numbers.")
    rec = []
    texts = {
        "Abitibi": ("Anchor hub region: 2 integrated mass-timber hubs, about 3 Indigenous-majority spokes, 2 bioheat "
                    "plants, 1 residue-based pellet plant.",
                    "Largest softwood surplus (about 440,000 m3/yr) and largest absolute gains in value added and Indigenous "
                    "income. Engage Treaty 9 First Nations in the Cochrane-Iroquois Falls area as equity partners in the hubs."),
        "TroutLake": ("Hub plus residue economy: 1 hub, 1-2 spokes sized to northern housing, 2 bioheat plants, 1 pellet plant.",
                      "Already the most utilized FMU (69%), reaching about 82%. Remote communities give the highest bioheat "
                      "value; residues from hub and incumbent mills feed bioheat and pellets."),
        "Kenora": ("Indigenous-led spokes first, hub conditional: up to 3 spokes and 2 bioheat plants; 1 hub only with "
                   "secured offtake.",
                   "Largest relative gain (15% to about 50%) and highest Indigenous share of benefits, but the highest risk "
                   "(hub and spoke failures) because of the roughly 1,900 km distance to the GTA. Markets in Manitoba and "
                   "north-western Ontario, not modelled here, would improve viability. Build on existing First Nations-led "
                   "forest management."),
        "Nipissing": ("Early hub plus full spoke network: 1-2 hubs, 3 spokes, 1 bioheat plant.",
                      "Close to southern markets with low freight. Utilization roughly doubles (31% to about 60%); strong "
                      "hub profitability makes it a priority first-wave location."),
        "Algoma": ("Hub plus spokes, hardwood pathway needed: 1 hub, 2-3 spokes, 1 bioheat plant.",
                   "Utilization rises from 20% to about 46%, limited by hardwood-dominated supply (tolerant hardwoods and "
                   "birch about 309,000 m3/yr LTMD) for which no new pathway is modelled."),
        "Algonquin": ("One hub and a smaller spoke network: about 1 hub, 2 spokes, 1 bioheat plant.",
                      "High hub profitability and short haul to markets, but limited free softwood. Further gains depend on "
                      "hardwood and white-pine products (glulam, hardwood CLT) and on Algonquins of Ontario participation."),
    }
    for r in ["Abitibi", "Nipissing", "Algoma", "TroutLake", "Algonquin", "Kenora"]:
        b, a = reg.loc[("S6", r)], reg.loc[("S0", r)]
        rec.append([REG[r], texts[r][0], f"{pct(cal.loc[r, 'utilization'])} -> {pct(b.util_late)}",
                    f"+{f0(b.rural_jobs_late - a.rural_jobs_late)} / +{f0(b.ind_jobs_late - a.ind_jobs_late)}",
                    f"+{f0(b.ind_income_npv - a.ind_income_npv)}", texts[r][1]])
    R.table(pd.DataFrame(rec, columns=["FMU", "Optimal portfolio", "Utilization", "Jobs: rural / Indigenous (FTE)",
                                       "Indigenous income ($M NPV)", "Rationale and conditions"]),
            "Optimal biomass utilization strategy by FMU region (S6 vs business as usual, base demand).",
            [2.0, 4.2, 1.8, 2.0, 1.8, 5.2], font=7.5)
    R.p("Cross-cutting elements of the optimal strategy", bold_lead="")
    R.bullets([
        "**A common provincial enabling framework. **Indigenous fibre-allocation priority on uncommitted Crown volume, "
        "capacity building for Indigenous harvesting contractors, procurement of mass timber and prefab panels in public and "
        "community housing, and a clean-heat incentive for bioheat. For less than $5 M of extra public cost, these "
        "instruments improve every outcome and reduce plant failures.",
        "**Indigenous ownership where the asset is community-scale; partnership where it is large. **Majority Indigenous "
        "ownership of spokes and bioheat plants with 40-50% capital grants, and around 20% equity partnerships in hubs.",
        "**Investment phasing by risk. **First wave: Abitibi River, Nipissing and Algoma hubs, plus bioheat and spokes in all "
        "FMUs. Second wave: Trout Lake and Algonquin Park hubs and pellets. Third wave: a Kenora hub only once offtake "
        "agreements are in place.",
        "**Size spokes to secured demand. **Spokes that rely on long-haul southern sales are the main source of stranded "
        "assets in remote FMUs; northern housing programs and procurement should anchor them.",
        "**Do not rely on urban mass-timber plants for rural or Indigenous objectives. **They are commercially viable, but "
        "they do not mobilize additional biomass.",
    ])

    # ---- 7 limitations ------------------------------------------------------------------------
    R.h("7. Limitations and next steps")
    R.bullets([
        "**Assumed parameters. **Several community-level inputs are planning assumptions that need validation with "
        "communities and MNR (Appendix A): northern housing need, community heat demand and fuel costs, Indigenous labour and "
        "contractor shares, and mass-timber and panel prices. They will be refined through the Activity 5 market survey and "
        "stakeholder consultations, including engagement with First Nations whose territories overlap the study FMUs.",
        "**No cross-FMU log trade. **Logs are sourced within each FMU and products are shipped to the GTA. Cross-FMU log trade, "
        "Manitoba/Prairie markets for Kenora and Trout Lake, and export markets for mass timber are not represented, so "
        "Kenora outcomes are likely conservative.",
        "**Missing poplar and hardwood pathway. **Poplar OSB/LVL and hardwood EWP options are not modelled; adding them is the "
        "priority model extension for Activity 6.",
        "**Not yet coupled to FORBIOM-ON. **The ABM shares data and parameters with FORBIOM-ON but is not dynamically coupled. "
        "The planned co-evolutionary link, with ABM adoption rates and harvest intensities passed to the GAMS model and shadow "
        "prices returned to agents, is the next methodological step.",
        "**Carbon accounting is partial. **It covers product storage, avoided fossil heat and process/transport emissions. "
        "Forest carbon dynamics and construction-material substitution are addressed in the Activity 4 life cycle assessment.",
        "**Indicative scale. **Results describe six FMUs, not all of Ontario. Absolute values should be read as indicative and "
        "relative rankings as robust.",
    ])

    # ---- references --------------------------------------------------------------------------
    R.h("References")
    for ref in [
        "Grimm, V., Railsback, S.F., Vincenot, C.E., et al. (2020). The ODD protocol for describing agent-based and other simulation models: a second update. Journal of Artificial Societies and Social Simulation 23(2): 7.",
        "Ontario Ministry of Natural Resources. Annual Report on Wood Supply (ARWoodS) data extract, 2025 update.",
        "Ontario Ministry of Natural Resources. Electronic Forest industry Annual Report (eFAR) summary data, 2020-2022 (September 2024 release).",
        "Ontario Ministry of Natural Resources and Forestry (2022). Forest Biomass Action Plan.",
        "Yousefpour, R., Elfahli, Y. (2026). Ontario Forest Biomass: Supply Chain Feasibility, Processing Optimization, and ABM-Integrated Biomass Utilization - preliminary results (Activities 1-3 interim report), University of Toronto.",
        "Yousefpour, R., Elfahli, Y. (2026). Future Pathways for Mass Timber and Offsite Wood Construction in Ontario: an economic evaluation and policy scan informed by the FORBIOM-ON model. FEDSS Lab, University of Toronto.",
        "Simon, H.A. (1957). Models of Man: Social and Rational. Wiley, New York.",
    ]:
        para = d.add_paragraph(ref)
        para.paragraph_format.left_indent = Cm(0.6)
        para.paragraph_format.first_line_indent = Cm(-0.6)
        para.runs[0].font.size = Pt(9)

    # ---- appendices --------------------------------------------------------------------------
    R.page_break()
    R.h("Appendix A. Parameters, sources and status")
    pr, tp = P["prices"], P["tech"]
    par = [
        ["LTMD and harvest by FMU and species group", "Table 2", "ARWoodS 2025, FY2019-2024 mean", "Calibrated"],
        ["Direct jobs per 1,000 m3, facility + woodlands (sawmill, composite)", f"{tp['inc_SW']['jobs']:.2f}, {tp['inc_PO']['jobs']:.2f}", "eFAR 2020-2022", "Calibrated"],
        ["Harvest cost (stump to roadside)", f"${P['harvest_cost']:.0f}/m3 x access factor 1.00-1.15", "FORBIOM-ON harvestCostRW", "Model-consistent"],
        ["In-region haul; stumpage", f"{P['inregion_haul_km']:.0f} km at ${P['freight_log']}/m3-km; ${P['stumpage']:.0f}/m3", "Ontario industry norms", "Literature/assumption"],
        ["Lumber (SPF) mill-net price", f"${pr['lumber']:.0f}/m3, AR(1) rho 0.6, sd 15%", "2023-2026 market range", "Literature/assumption"],
        ["CLT/glulam price", f"${pr['mt']:.0f}/m3 (sensitivity +/-20%)", "Industry quotes; FORBIOM-ON ratio tested", "Assumption, to validate (Activity 5)"],
        ["Prefab panel price (wood-equivalent)", f"${pr['panel']:.0f}/m3", "Industry estimates", "Assumption, to validate (Activity 5)"],
        ["OSB; pulpwood; mill residue; pellets", f"${pr['osb']:.0f}/m3; ${pr['pulpwood']:.0f}/m3; ${pr['residue']:.0f}/odt; ${pr['pellet']:.0f}/odt", "Market references", "Literature/assumption"],
        ["Lumber recovery; residue fraction (sawmill)", f"{tp['inc_SW']['lrf']}; {tp['inc_SW']['res_frac']}", "Boreal sawmill literature", "Literature"],
        ["Lumber per m3 CLT / panel", f"{tp['mt_hub']['lumber_per_mt']} / {tp['spoke']['lumber_per_panel']}", "FORBIOM-ON lumberInputPerUnit", "Model-consistent"],
        ["Mass-timber hub", f"150k m3 logs, 60k m3 MT, capex $150M, MT conversion ${tp['mt_hub']['mt_conv']:.0f}/m3", "Comparable plants", "Assumption"],
        ["Community spoke", f"25k m3 logs, 10k m3 panels, capex $25M, panel conversion ${tp['spoke']['panel_conv']:.0f}/m3, {tp['spoke']['jobs_panel']:.0f} jobs per 1,000 m3 panels", "Modular plant benchmarks", "Assumption"],
        ["Bioheat; pellet plant", f"5k odt, $12M; 80k odt, $50M", "Project benchmarks", "Assumption"],
        ["Mass-timber demand (six regions)", f"{P['mt_demand0']:,.0f} m3/yr in 2026; growth 3/8/15%", "FORBIOM-ON demand caps", "Model-consistent / scenario"],
        ["Southern prefab panel demand", f"{P['panel_demand0']:,.0f} m3/yr", "Planning assumption", "To validate (Activity 5)"],
        ["Northern housing need (units/yr)", ", ".join(f"{REG[r]} {v}" for r, v in P["north_housing_units"].items()), "Planning assumption", "To validate with communities"],
        ["Community heat demand (odt/yr)", ", ".join(f"{REG[r]} {v:,}" for r, v in P["heat_demand_odt"].items()), "Planning assumption", "To validate with communities"],
        ["Displaced heat value ($/odt)", ", ".join(f"{REG[r]} {v}" for r, v in P["heat_value_odt"].items()), "Heating oil/propane/diesel costs", "Assumption"],
        ["Indigenous labour share", ", ".join(f"{REG[r]} {v:.0%}" for r, v in P["ind_labour_share"].items()), "Approx. Census 2021 district shares", "To verify"],
        ["Indigenous contractor share (initial)", ", ".join(f"{REG[r]} {v:.0%}" for r, v in P["ind_contractor_share"].items()), "Planning assumption", "To validate with MNR/communities"],
        ["Discount rate (private; social)", "8%; 4%", "Standard practice", "Sensitivity 10%"],
        ["Carbon factors", "FORBIOM-ON storage/process factors; 1.23 tCO2e/odt avoided by bioheat", "FORBIOM-ON; fuel emission factors", "Model-consistent"],
        ["Behaviour", "Contractor k=12; mill k=4; investment k=5, max probability 0.6; expectation weight 0.3; price adjustment 0.3", "Calibrated for plausible response", "Tested in sensitivity"],
    ]
    R.table(pd.DataFrame(par, columns=["Parameter", "Value", "Source / basis", "Status"]),
            "Model parameters (2026 CAD, real). Values marked 'to validate' are planning assumptions to be refined through "
            "Activity 5 consultations.", [4.6, 5.0, 4.0, 3.0], font=7.5)

    R.h("Appendix B. Software and reproducibility")
    R.p("FORBIOM-ABM is implemented in Python 3.11 (NumPy, pandas, Matplotlib) and runs on a standard laptop. "
        "Files in the ABM_Activity3 folder:")
    R.bullets([
        "calibrate.py: builds calibration tables from ARWoodS and eFAR.",
        "abm.py: model (agents, markets, strategies, parameters).",
        "experiments.py: Monte Carlo runs, regional tuning (S6), sensitivity analysis and MCDA.",
        "analysis.py, figures.py, build_report.py: tables, figures and this report.",
        "To reproduce: py calibrate.py; py experiments.py --seeds 50 --sens-seeds 30; py analysis.py; py figures.py; "
        "py build_report.py.",
    ])

    R.h("Appendix C. Region-by-region tuning of the enabling-policy hybrid (S5)")
    vt = variants.copy()
    vt["_r"] = vt.region.map(REGIONS.index)
    vt["_v"] = (vt.variant != "full").astype(int)
    vt = vt.sort_values(["_r", "_v", "variant"])
    vt["FMU"] = vt.region.map(REG)
    vt = vt[["FMU", "variant", "util_late", "va_npv", "rural_jobs_late", "ind_income_npv", "public_cost", "risk_unprofitable", "gain_regional", "gain_provincial"]]
    vt["util_late"] = vt.util_late.map(pct)
    for c in ("va_npv", "rural_jobs_late", "ind_income_npv", "public_cost"):
        vt[c] = vt[c].map(f0)
    vt["risk_unprofitable"] = vt.risk_unprofitable.map(pct)
    vt["gain_regional"] = vt.gain_regional.map(lambda x: f"{x:+.3f}")
    vt["gain_provincial"] = vt.gain_provincial.map(lambda x: f"{x:+.3f}")
    vt.columns = ["FMU", "Variant", "Utilization", "VA NPV ($M)", "Rural jobs", "Indig. income ($M)", "Public cost ($M)",
                  "Unprofitable", "Regional gain", "Provincial gain"]
    R.table(vt, "Effect of removing one technology (or capping spokes at one) in one FMU, relative to full S5. Gains are "
            "weighted changes in units of the S0-to-S5 improvement; a variant is adopted only if it is no worse regionally "
            "and provincially and reduces public cost. Selected: " + "; ".join(f"{REG[r]}: {v}" for r, v in s6def["regional_choice"].items()) + ".",
            font=7)

    d.save(OUT)
    print("saved", OUT)

    # ---- Appendix I (interim report) text for Activity 3 ---------------------------------------
    A = Document()
    A.styles["Normal"].font.name = "Calibri"
    A.styles["Normal"].font.size = Pt(10.5)
    A.add_heading("Appendix I - Activity 3 \"Detailed Results\" (replacement text)", level=2)
    items = [
        "Status and variation: Completed October 2026 (milestone date March 31, 2026). By March 31 the data compilation and "
        "ABM framework were in place, as presented in the preliminary results report. The additional time was used to calibrate "
        "the model to observed FMU-level harvest and employment, align it with the FORBIOM-ON model developed under "
        "Activity 2, and run the full simulation experiments. The milestone deliverable, the ABM simulation report identifying "
        "optimal biomass utilization strategies, is attached.",
        "Compiled and structured FMU-level fibre supply data (LTMD and actual harvest by species group, ARWoodS, 2019-2024), "
        "facility employment and wood-flow data (eFAR, 2020-2022), and technical, cost, demand and carbon parameters shared "
        "with FORBIOM-ON (Activity 2), for six study FMUs: Trout Lake, Abitibi River, Algoma, Kenora, Nipissing and Algonquin Park.",
        "Developed FORBIOM-ABM, an agent-based model representing FMU managers, harvest contractors (Indigenous-owned and "
        "other), incumbent mills, new processing investors (northern mass-timber hubs, urban CLT plants, community "
        "log-to-panel plants, community bioheat, pellet plants), product and heat markets, and a policy agent. Agents interact "
        "annually from 2026 to 2045 through regional fibre markets with endogenous log prices and through southern "
        "mass-timber and prefab markets.",
        "Implemented behavioural decision rules: logistic (bounded-rationality) supply, capacity-utilization and investment "
        "rules; adaptive price expectations; NPV-based investment with learning-by-doing and build lags; exit after sustained "
        "losses; and FMU allocation rules including an Indigenous allocation-priority option. Indigenous employment, "
        "contractor income and equity returns are tracked explicitly.",
        f"Calibrated the model to observed harvest. Simulated start-year harvest matches ARWoodS within {max_err:.1f}% in "
        f"every FMU, and business-as-usual runs reproduce the persistent gap between harvest and LTMD "
        f"({pct(act_tot / ltmd_tot)} utilization across the six FMUs).",
        f"Simulated seven biomass utilization strategies, combining technology options and policy instruments, with {n_runs} "
        "Monte Carlo replications under three mass-timber demand scenarios and six sensitivity cases. Strategies were "
        "evaluated with a seven-criterion multi-criteria analysis: utilization, value added, rural jobs, Indigenous income, "
        "carbon, public cost and investment risk.",
        f"Results: the optimal strategy is a regionally differentiated hub-and-spoke portfolio. It combines northern "
        f"log-to-mass-timber hubs, Indigenous-majority community panel and bioheat plants, and pellet plants in Abitibi River "
        f"and Trout Lake only, backed by Indigenous allocation priority, contractor capacity building and public procurement. "
        f"Relative to business as usual it raises utilization from {pct(s0.util_late)} to {pct(s6.util_late)}, rural jobs from "
        f"{f0(s0.rural_jobs_late)} to {f0(s6.rural_jobs_late)}, Indigenous jobs from {f0(s0.ind_jobs_late)} to "
        f"{f0(s6.ind_jobs_late)}, Indigenous community income from ${f0(s0.ind_income_npv)} M to ${f0(s6.ind_income_npv)} M "
        f"(NPV) and value added by ${f0(va_gain)} M (NPV), for about ${f0(s6.public_cost)} M in public support. It ranked first "
        "under all demand scenarios, sensitivity cases and stakeholder weightings.",
        "Identified region-specific strategies: Abitibi River as the anchor hub region; early hubs and full spoke networks in "
        "Nipissing and Algoma; a hub plus residue-based bioheat and pellets in Trout Lake; a smaller portfolio in Algonquin Park; "
        "and Indigenous-led spokes first in Kenora, with a hub conditional on secured offtake because of distance-driven risk.",
        "Key outputs feed Activities 4-6: carbon results for the LCA; community-level assumptions (housing need, heat demand, "
        "Indigenous participation) to be validated through the Activity 5 consultations; and the poplar/hardwood utilization "
        "gap as a priority for the Activity 6 policy recommendations.",
    ]
    for it in items:
        A.add_paragraph(it, style="List Bullet")
    out_a = HERE / "Appendix_I_Activity3_DetailedResults.docx"
    A.save(out_a)
    print("saved", out_a)


if __name__ == "__main__":
    main()
