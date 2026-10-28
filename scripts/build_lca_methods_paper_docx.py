"""Build a short technical methods-and-data paper (DOCX) from the LCA assessment."""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor, Cm

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
MD = os.path.join(ROOT, "docs", "LCA_Economic_Impact_Methods_and_Data.md")
OUT = os.path.join(ROOT, "docs", "LCA_Economic_Impact_Methods_and_Data.docx")

FOREST, INK, GREY = "1B4332", "1F2A24", "5F6B64"
BODY, HEAD = "Calibri", "Georgia"


def rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_font(run, name):
    run.font.name = name
    rFonts = run._r.get_or_add_rPr().find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        run._r.get_or_add_rPr().append(rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(a), name)


def run_fmt(run, size=11, color=INK, bold=False, italic=False, font=BODY):
    set_font(run, font)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.4)
    s.right_margin = Cm(2.4)

normal = doc.styles["Normal"]
normal.font.name = BODY
normal.font.size = Pt(11)
normal.font.color.rgb = rgb(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

text = open(MD, encoding="utf-8").read().splitlines()
i = 0
while i < len(text):
    line = text[i]
    if not line.strip():
        i += 1
        continue
    if line.startswith("# "):
        p = doc.add_paragraph()
        r = p.add_run(line[2:].strip())
        run_fmt(r, size=16, color=FOREST, bold=True, font=HEAD)
        p.paragraph_format.space_after = Pt(4)
    elif line.startswith("## "):
        p = doc.add_paragraph()
        r = p.add_run(line[3:].strip())
        run_fmt(r, size=13, color=FOREST, bold=True, font=HEAD)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
    elif line.startswith("### "):
        p = doc.add_paragraph()
        r = p.add_run(line[4:].strip())
        run_fmt(r, size=11.5, color=FOREST, bold=True, italic=True, font=HEAD)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
    elif line.startswith("---"):
        i += 1
        continue
    elif line.startswith("|") and i + 1 < len(text) and re.match(r"^\|[-:| ]+\|$", text[i + 1].strip()):
        rows = []
        while i < len(text) and text[i].startswith("|"):
            if re.match(r"^\|[-:| ]+\|$", text[i].strip()):
                i += 1
                continue
            cells = [c.strip() for c in text[i].strip().strip("|").split("|")]
            rows.append(cells)
            i += 1
        if not rows:
            continue
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                cell.text = ""
                p = cell.paragraphs[0]
                r = p.add_run(val)
                run_fmt(r, size=9, color=INK, bold=(ri == 0))
        doc.add_paragraph()
        continue
    elif line.startswith("- ") or re.match(r"^\d+\. ", line):
        p = doc.add_paragraph(style="List Bullet" if line.startswith("- ") else "List Number")
        content = re.sub(r"^(\- |\d+\. )", "", line)
        # bold **segments**
        parts = re.split(r"(\*\*.+?\*\*)", content)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                run_fmt(r, bold=True)
            else:
                r = p.add_run(part)
                run_fmt(r)
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        parts = re.split(r"(\*\*.+?\*\*|https?://\S+)", line)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                run_fmt(r, bold=True)
            elif part.startswith("http"):
                r = p.add_run(part)
                run_fmt(r, color="2D6A4F", italic=True)
            else:
                r = p.add_run(part)
                run_fmt(r, size=11 if not line.startswith("**Rasoul") else 11)
        if line.startswith("**Rasoul") or line.startswith("University") or line.startswith("October"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                run_fmt(r, size=10, color=GREY, italic=True)
    i += 1

doc.core_properties.title = "Methods and data for the life-cycle and economic-impact assessment"
doc.core_properties.author = "Rasoul Yousefpour; Yasser Elfahli"
doc.save(OUT)
print("Saved", OUT)
